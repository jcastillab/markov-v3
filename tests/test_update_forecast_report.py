import numpy as np
import pandas as pd
import zipfile
from openpyxl import load_workbook

from scripts.package_portable import build_zip
from scripts.update_forecast_report import _build_reports, _write_report_workbook


def test_report_keeps_future_week_without_inventing_real_values():
    weekly = pd.DataFrame([{
        "modelo": "RF_OPT_FENO_ENTRADA",
        "finca": "ALMER",
        "bloque": "1",
        "fecha_origen": "2026-09-21",
        "semana_proyeccion": 202640,
        "real": np.nan,
        "proyectado": 1234.5,
        "dias": 7,
        "dias_reales": 0,
        "estado_evaluacion": "NO EVALUABLE",
    }])

    farm, block = _build_reports(weekly)

    assert farm.loc[0, "semana_proyeccion"] == 202640
    assert block.loc[0, "semana_proyeccion"] == 202640
    assert farm.loc[0, "estado_evaluacion"] == "NO EVALUABLE"
    assert block.loc[0, "estado_evaluacion"] == "NO EVALUABLE"
    assert pd.isna(farm.loc[0, "real"])
    assert pd.isna(block.loc[0, "real"])
    assert farm.loc[0, "proyectado"] == 1234.5
    assert farm.loc[0, "indicador"] == "SIN REAL"


def test_report_workbook_is_written_without_excel_automation(tmp_path):
    path = tmp_path / "report.xlsx"
    _write_report_workbook(path, {
        "COMPARACION_SEMANAL": pd.DataFrame({"semana": [202640], "real": [None]}),
        "COMPARACION_SEMANAL_BLOQUE": pd.DataFrame({"bloque": ["1"], "proyectado": [4.5]}),
    })

    workbook = load_workbook(path, data_only=True)
    assert workbook.sheetnames == ["COMPARACION_SEMANAL", "COMPARACION_SEMANAL_BLOQUE"]
    assert workbook["COMPARACION_SEMANAL"]["A2"].value == 202640
    assert workbook["COMPARACION_SEMANAL"]["B2"].value is None
    assert workbook["COMPARACION_SEMANAL"].tables["Table1"].ref == "A1:B2"
    assert workbook["COMPARACION_SEMANAL_BLOQUE"]["B2"].value == 4.5


def test_portable_package_does_not_contain_powershell_files(tmp_path):
    project = tmp_path / "project"
    (project / "scripts").mkdir(parents=True)
    (project / "scripts" / "entry.py").write_text("print('ok')", encoding="utf-8")
    output = tmp_path / "portable.zip"

    build_zip(project, output)

    with zipfile.ZipFile(output) as package:
        names = package.namelist()
        assert not any(name.lower().endswith((".ps1", ".psm1", ".psd1")) for name in names)
        readme = package.read("markov_v3/LEEME_PRIMERO.txt").decode("utf-8")
        assert "powershell" not in readme.lower()
