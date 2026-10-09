"""Genera un informe Excel de proyecciones con Python puro (openpyxl)."""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
MODELS = {
    "RF_H1_H7_FENO_ENTRADA",
    "RF_OPT_FENO_ENTRADA",
    "GLM_NB_FENO_CLIMA_ENTRADA",
    "NB_JERARQUICO_COVARIABLES_ENTRADA",
}


def _write_report_workbook(path: Path, sources: dict[str, pd.DataFrame]) -> None:
    """Escribe ambas tablas y formato básico sin Excel, COM ni plantilla propietaria."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, (sheet_name, frame) in enumerate(sources.items(), start=1):
        sheet = workbook.create_sheet(sheet_name)
        headers = list(frame.columns)
        sheet.append(headers)
        for row in frame.itertuples(index=False, name=None):
            sheet.append([
                None if pd.isna(value) else
                value.item() if hasattr(value, "item") else value
                for value in row
            ])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        if headers:
            table = Table(displayName=f"Table{index}", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False,
            )
            sheet.add_table(table)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                values = [len(str(cell.value or "")) for cell in column]
                sheet.column_dimensions[column[0].column_letter].width = min(
                    max(max(values, default=0) + 2, 12), 34
                )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _build_reports(weekly: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Produce las dos formas tabulares que alimentan las tablas dinámicas."""
    if {"modelo", "finca", "bloque", "fecha_origen", "semana_proyeccion",
        "real", "proyectado"}.difference(weekly.columns):
        raise ValueError("El CSV semanal de entrada no contiene las columnas requeridas")
    weekly = weekly[weekly["modelo"].isin(MODELS)].copy()
    weekly["semana_proyeccion"] = pd.to_numeric(
        weekly["semana_proyeccion"], errors="coerce").astype("Int64")
    weekly["fecha_origen"] = pd.to_datetime(weekly["fecha_origen"], errors="coerce")
    if weekly.empty:
        raise ValueError("No hay filas de los modelos del informe en la evaluación actual")
    # Los informes establecidos seleccionan el origen representativo por
    # modelo/bloque/semana; las filas ya vienen agregadas y se evalúan aquí.
    from src.evaluacion_entrada import _status, _select_weekly_entry_origins
    selected = _select_weekly_entry_origins(weekly)
    selected["diferencia"] = selected["proyectado"] - selected["real"]
    selected["razon_proyectado_real"] = selected["proyectado"] / selected["real"].where(selected["real"].ne(0))
    selected["desviacion_pct"] = selected["razon_proyectado_real"] - 1
    selected["indicador"] = selected["razon_proyectado_real"].map(_status)
    selected.loc[selected["estado_evaluacion"].eq("NO EVALUABLE"),
                 ["real", "diferencia", "desviacion_pct", "indicador"]] = [pd.NA, pd.NA, pd.NA, "SIN REAL"]
    selected["modelo_nombre"] = selected["modelo"].map({
        "RF_H1_H7_FENO_ENTRADA": "RF H1-H7 FENO",
        "RF_OPT_FENO_ENTRADA": "RF global OPT",
        "GLM_NB_FENO_CLIMA_ENTRADA": "GLM NB FENO + clima",
        "NB_JERARQUICO_COVARIABLES_ENTRADA": "Bayes NB - covariables",
    })
    block = selected[["modelo", "modelo_nombre", "finca", "bloque", "semana_proyeccion",
                      "estado_evaluacion", "real", "proyectado", "diferencia",
                      "desviacion_pct", "indicador"]].copy()

    farm = (selected.groupby(["modelo", "finca", "semana_proyeccion"], as_index=False,
                             dropna=False)
            .agg(real=("real", "sum"), proyectado=("proyectado", "sum"),
                 dias=("dias", "sum"), dias_reales=("dias_reales", "sum")))
    farm["estado_evaluacion"] = "NO EVALUABLE"
    farm.loc[farm.dias_reales.eq(farm.dias), "estado_evaluacion"] = "COMPLETA"
    farm.loc[farm.dias_reales.gt(0) & farm.dias_reales.lt(farm.dias), "estado_evaluacion"] = "PARCIAL"
    farm["diferencia"] = farm["proyectado"] - farm["real"]
    farm["razon_proyectado_real"] = farm["proyectado"] / farm["real"].where(farm["real"].ne(0))
    farm["desviacion_pct"] = farm["razon_proyectado_real"] - 1
    farm["indicador"] = farm["razon_proyectado_real"].map(_status)
    farm.loc[farm.estado_evaluacion.eq("NO EVALUABLE"),
             ["real", "diferencia", "desviacion_pct", "indicador"]] = [pd.NA, pd.NA, pd.NA, "SIN REAL"]
    farm["modelo_nombre"] = farm["modelo"].map(selected.drop_duplicates("modelo").set_index("modelo")["modelo_nombre"])
    farm = farm[["modelo", "modelo_nombre", "finca", "semana_proyeccion", "estado_evaluacion",
                 "real", "proyectado", "diferencia", "desviacion_pct", "indicador"]]
    return farm, block


def run(root: Path = ROOT) -> Path:
    from src.canonical import load_config, pipeline_config_path

    cfg = load_config(pipeline_config_path(root))
    profile_outputs = root / cfg["paths"]["outputs"]
    weekly_csv = profile_outputs / "evaluation" / "predictions_entrada_semanales.csv"
    output = profile_outputs / "informe pronosticos.xlsx"
    if not weekly_csv.is_file():
        raise FileNotFoundError(f"No se encontró evaluación del dashboard: {weekly_csv}")
    weekly = pd.read_csv(weekly_csv, parse_dates=["fecha_origen"])
    farm, block = _build_reports(weekly)
    maximum_week = int(pd.to_numeric(block["semana_proyeccion"]).max())
    _write_report_workbook(output, {
        "COMPARACION_SEMANAL": farm,
        "COMPARACION_SEMANAL_BLOQUE": block,
    })
    print(f"Informe actualizado: {output}")
    print(f"Última semana incluida: {maximum_week}")
    print(f"Filas por finca: {len(farm)}; por finca-bloque: {len(block)}")
    return output


if __name__ == "__main__":
    try:
        run()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
