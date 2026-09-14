"""Evalua modelos entrenados con el historico sobre una entrada externa.

El archivo externo solo aporta features en el origen y reales para scoring.
Los estimadores directos se ajustan exclusivamente con el dataset historico
original, nunca con ``Cantidad`` de la entrada evaluada.
"""

from __future__ import annotations

import json
import hashlib
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

try:
    from canonical import (active_beds_by_date, build_forecast_windows,
                           load_bed_validity, load_config, load_sampled_beds)
    from evaluation.metrics import metrics
    from models.bayes import CovariateHierarchicalNB, DirichletM3, HierarchicalNB
    from models.m3 import M3Matrix, _period_for_date, fit_m3, load_traditional_intervals, simulate
    from models.selection import estimator_from_spec
    from models.supervised import NegativeBinomialGLM, build_supervised_dataset, feature_groups
    from modeling.extrapolation import attach_extrapolation
except ModuleNotFoundError:
    from src.canonical import (active_beds_by_date, build_forecast_windows,
                               load_bed_validity, load_config, load_sampled_beds)
    from src.evaluation.metrics import metrics
    from src.models.bayes import CovariateHierarchicalNB, DirichletM3, HierarchicalNB
    from src.models.m3 import M3Matrix, _period_for_date, fit_m3, load_traditional_intervals, simulate
    from src.models.selection import estimator_from_spec
    from src.models.supervised import NegativeBinomialGLM, build_supervised_dataset, feature_groups
    from src.modeling.extrapolation import attach_extrapolation


MODEL_NAMES = (
    "E00_M3_BASE_ENTRADA",
    "E01_M3_INGRESO_CALIBRADO_ENTRADA",
    "RF_OPT_FENO_ENTRADA",
    "RF_OPT_FENO_PODA_ENTRADA",
    "RF_OPT_FENO_CLIMA_ENTRADA",
    "RF_OPT_FENO_PODA_CLIMA_ENTRADA",
    "GLM_NB_FENO_ENTRADA",
    "GLM_NB_FENO_PODA_ENTRADA",
    "GLM_NB_FENO_CLIMA_ENTRADA",
    "GLM_NB_FENO_PODA_CLIMA_ENTRADA",
    "RF_RESIDUAL_M3_FENO_ENTRADA",
    "RF_H1_H7_FENO_ENTRADA",
    "NB_JERARQUICO_ENTRADA",
    "NB_JERARQUICO_COVARIABLES_ENTRADA",
    "M3_DIRICHLET_MULTINOMIAL_ENTRADA",
)


def _read_input(path: Path, cfg: dict, template: pd.DataFrame, raw: Path) -> pd.DataFrame:
    source = pd.read_excel(path)
    required = {"Finca", "Bloque", "Fecha", "Cantidad", "semana", "conteo_RC",
                "conteo_SS", "conteo_AP", "conteo_CO", "conteo_total"}
    missing = required.difference(source.columns)
    if missing:
        raise ValueError(f"Faltan columnas en la entrada: {sorted(missing)}")
    aliases = cfg["farm_aliases"]
    out = pd.DataFrame({
        "finca": source["Finca"].map(lambda value: _canonical_farm(value, aliases)),
        "bloque": source["Bloque"].map(_canonical_block),
        "fecha": pd.to_datetime(source["Fecha"], dayfirst=True, errors="coerce").dt.normalize(),
        "semana_iso": pd.to_numeric(source["semana"], errors="coerce").astype("Int64"),
        "corte_comercial_real": pd.to_numeric(source["Cantidad"], errors="coerce"),
        "conteo_RC": pd.to_numeric(source["conteo_RC"], errors="coerce"),
        "conteo_SS": pd.to_numeric(source["conteo_SS"], errors="coerce"),
        "conteo_AP": pd.to_numeric(source["conteo_AP"], errors="coerce"),
        "conteo_CO": pd.to_numeric(source["conteo_CO"], errors="coerce"),
        "conteo_total": pd.to_numeric(source["conteo_total"], errors="coerce"),
    })
    out = out[out.finca.notna() & out.fecha.notna()].copy()
    if out.duplicated(["finca", "bloque", "fecha"]).any():
        raise ValueError("La entrada tiene claves finca+bloque+fecha duplicadas")
    out["es_fecha_conteo"] = out["conteo_total"].notna()

    # Las dimensiones estaticas y las features auxiliares ya auditadas se
    # reutilizan por clave; los conteos y el real siempre vienen de la entrada.
    # La escala de muestreo se recalcula abajo: nunca se hereda por fecha desde
    # el historico, porque una entrada puede contener semanas nuevas.
    aux = [c for c in template.columns if c not in {
        "finca", "bloque", "fecha", "semana_iso", "corte_comercial_real",
        "conteo_RC", "conteo_SS", "conteo_AP", "conteo_CO", "conteo_total",
        "es_fecha_conteo", "fecha_conteo_origen", "dias_desde_conteo",
        "conteo_RC_origen", "conteo_SS_origen", "conteo_AP_origen",
        "conteo_CO_origen", "conteo_total_origen", "semana_objetivo",
        "camas_activas", "camas_muestreadas", "factor_extrapolacion",
        "cobertura_muestreo", "plantas_activas", "area_activa",
    }]
    lookup = template[["finca", "bloque", "fecha"] + aux].drop_duplicates(
        ["finca", "bloque", "fecha"])
    out = out.merge(lookup, on=["finca", "bloque", "fecha"], how="left")
    defaults = {"poda_alineamiento": 0.0, "poda_corte": 0.0, "poda_total": 0.0}
    for col, default in defaults.items():
        if col in out:
            out[col] = out[col].fillna(default)

sampled = load_sampled_beds(raw, cfg)
    beds = load_bed_validity(raw, cfg)
    active = active_beds_by_date(beds, out["fecha"])
    out = attach_extrapolation(out, sampled, active, strict=False)
    origins = (out[out["es_fecha_conteo"]].sort_values("fecha")
               .groupby(["finca", "bloque", "semana_iso"], as_index=False).tail(1)
               [["finca", "bloque", "semana_iso", "fecha", "conteo_RC", "conteo_SS",
                 "conteo_AP", "conteo_CO", "conteo_total"]]
               .rename(columns={"fecha": "fecha_conteo_origen",
                                "conteo_RC": "conteo_RC_origen", "conteo_SS": "conteo_SS_origen",
                                "conteo_AP": "conteo_AP_origen", "conteo_CO": "conteo_CO_origen",
                                "conteo_total": "conteo_total_origen"}))
    out = out.merge(origins, on=["finca", "bloque", "semana_iso"], how="left", validate="many_to_one")
    out["dias_desde_conteo"] = (out["fecha"] - out["fecha_conteo_origen"]).dt.days
    out["semana_objetivo"] = out["semana_iso"] + 1
    return out


def _canonical_farm(value: object, aliases: dict) -> str | None:
    if pd.isna(value):
        return None
    key = " ".join(str(value).strip().upper().split())
    normalized = {" ".join(str(k).strip().upper().split()): v for k, v in aliases.items()}
    return normalized.get(key)


def _canonical_block(value: object) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


def _base_rows(frame: pd.DataFrame, mask: np.ndarray, model: str, pred) -> pd.DataFrame:
    cols = ["finca", "bloque", "fecha_origen", "fecha_objetivo", "semana_proyeccion",
            "horizonte_dia", "estado_ventana", "target"]
    out = frame.loc[mask, cols].rename(columns={"target": "real"}).copy()
    out["modelo"] = model
    out["proyectado"] = np.asarray(pred, float)
    for column in ("camas_activas", "camas_muestreadas", "factor_extrapolacion",
                   "cobertura_muestreo"):
        if column in frame:
            out[column] = frame.loc[mask, column].to_numpy()
    return out


def _m3(frame: pd.DataFrame, intervals: pd.DataFrame, cfg: dict, alpha: float, name: str) -> pd.DataFrame:
    rows = []
    matrices = {}
    for _, row in frame.iterrows():
        origin = pd.Timestamp(row.fecha_origen)
        period = _period_for_date(origin, cfg["m3"]["periods"])
        key = (row.finca, period, origin)
        if key not in matrices:
            matrices[key] = fit_m3(intervals, row.finca, period, origin)
        matrix = matrices[key]
        x0 = np.array([row.RC_t0, row.SS_t0, row.AP_t0], float)
        lead = (pd.Timestamp(row.fecha_objetivo) - origin).days
        pred = simulate(matrix, x0, lead, alpha).iloc[-1].PC_dia_muestra * row.factor_extrapolacion
        result = _base_rows(frame, frame.index == row.name, name, [pred])
        rows.append(result)
    return pd.concat(rows, ignore_index=True) if rows else pd.DataFrame()


def _direct_models(frame: pd.DataFrame, historical: pd.DataFrame, cfg: dict,
                   rf_specs: dict[str, dict]) -> dict[str, pd.DataFrame]:
    groups = feature_groups(frame)
    outputs = {}
    numeric = frame.select_dtypes(include=[np.number]).columns
    historical = historical[historical.target.notna()].copy()
    if historical.empty:
        raise ValueError("No hay historico original con target para entrenar")
    for name, group_name, estimator_factory in [
        (f"GLM_NB_{group}_ENTRADA", group, lambda: NegativeBinomialGLM(
            cfg["supervised"]["nb_alpha"], cfg["supervised"]["nb_max_iter"]))
        for group in groups
    ]:
        cols = list(dict.fromkeys(c for c in groups[group_name] if c in numeric and c != "target"))
        x_new = frame[cols].replace([np.inf, -np.inf], np.nan).fillna(0)
        estimator = estimator_factory().fit(
            historical[cols].replace([np.inf, -np.inf], np.nan).fillna(0), historical.target)
        outputs[name] = _base_rows(frame, np.ones(len(frame), dtype=bool), name, estimator.predict(x_new))

    for name, spec in rf_specs.items():
        group_name = spec["features"]
        cols = list(dict.fromkeys(c for c in groups[group_name] if c in numeric and c != "target"))
        x_new = frame[cols].replace([np.inf, -np.inf], np.nan).fillna(0)
        estimator = estimator_from_spec(spec, cfg).fit(
            historical[cols].replace([np.inf, -np.inf], np.nan).fillna(0), historical.target)
        outputs[name] = _base_rows(frame, np.ones(len(frame), dtype=bool), name, estimator.predict(x_new))
    return outputs


def _additional_rf_models(frame: pd.DataFrame, historical: pd.DataFrame, cfg: dict,
                          spec: dict) -> dict[str, pd.DataFrame]:
    groups = feature_groups(frame)
    cols = list(dict.fromkeys(c for c in groups["FENO"]
                              if c in frame.select_dtypes(include=[np.number]).columns and c != "target"))
    train = historical[historical.target.notna()].copy()
    if train.empty:
        return {"RF_RESIDUAL_M3_FENO_ENTRADA": pd.DataFrame(), "RF_H1_H7_FENO_ENTRADA": pd.DataFrame()}
    eligible = np.ones(len(frame), dtype=bool)
    x_train = train[cols].replace([np.inf, -np.inf], np.nan).fillna(0)
    x_new = frame[cols].replace([np.inf, -np.inf], np.nan).fillna(0)
    residual_target = train.target - train.M3_pred_bloque
    residual_valid = np.isfinite(residual_target.to_numpy())
    if not residual_valid.any():
        residual = pd.DataFrame()
    else:
        residual_estimator = estimator_from_spec(spec, cfg).fit(
            x_train.loc[residual_valid], residual_target.loc[residual_valid])
        residual_pred = np.maximum(0, frame.M3_pred_bloque.to_numpy() +
                                  residual_estimator.predict(x_new))
        residual = _base_rows(frame, eligible, "RF_RESIDUAL_M3_FENO_ENTRADA", residual_pred[eligible])
    horizon_rows = []
    for horizon in sorted(pd.to_numeric(train.horizonte_dia).dropna().unique()):
        train_h = train[train.horizonte_dia.eq(horizon)]
        current_h = eligible & frame.horizonte_dia.eq(horizon).to_numpy()
        if train_h.empty or not current_h.any():
            continue
        estimator = estimator_from_spec(spec, cfg).fit(
            train_h[cols].replace([np.inf, -np.inf], np.nan).fillna(0), train_h.target)
        horizon_rows.append(_base_rows(frame, current_h, "RF_H1_H7_FENO_ENTRADA",
                                       estimator.predict(x_new.loc[current_h])))
    return {"RF_RESIDUAL_M3_FENO_ENTRADA": residual,
            "RF_H1_H7_FENO_ENTRADA": pd.concat(horizon_rows, ignore_index=True) if horizon_rows else pd.DataFrame()}


def _bayes_models(frame: pd.DataFrame, historical: pd.DataFrame, intervals: pd.DataFrame,
                  cfg: dict) -> dict[str, pd.DataFrame]:
    outputs = {"NB_JERARQUICO_ENTRADA": [], "NB_JERARQUICO_COVARIABLES_ENTRADA": [],
               "M3_DIRICHLET_MULTINOMIAL_ENTRADA": []}
    covariates = ["RC_t0", "SS_t0", "AP_t0", "p_RC", "p_SS", "p_AP", "TOTAL_t0",
                  "factor_extrapolacion", "corte_lag_1d", "corte_lag_2d", "corte_lag_3d",
                  "corte_sum_3d", "corte_sum_7d", "corte_sum_14d", "corte_mean_7d",
                  "horizonte_dia"]
    train = historical[historical.target.notna()].copy()
    if train.empty:
        raise ValueError("No hay historico original con target para entrenar")
    nb = HierarchicalNB(cfg["bayes"]["hierarchical_shrinkage"]).fit(train)
    cov = CovariateHierarchicalNB(cfg["bayes"]["hierarchical_shrinkage"], cfg["bayes"]["covariate_ridge"])
    cov.fit(train, [c for c in covariates if c in train.columns])
    posterior_cache = {}
    for origin in sorted(pd.to_datetime(frame.fecha_origen).unique()):
        current = pd.to_datetime(frame.fecha_origen).eq(origin).to_numpy()
        current_frame = frame.loc[current]
        outputs["NB_JERARQUICO_ENTRADA"].append(_base_rows(
            frame, current, "NB_JERARQUICO_ENTRADA", nb.predict(current_frame)))
        outputs["NB_JERARQUICO_COVARIABLES_ENTRADA"].append(_base_rows(
            frame, current, "NB_JERARQUICO_COVARIABLES_ENTRADA", cov.predict(current_frame)))
        predictions = []
        for _, row in current_frame.iterrows():
            period = _period_for_date(origin, cfg["m3"]["periods"])
            cache_key = (row.finca, period, origin)
            if cache_key not in posterior_cache:
                prior = fit_m3(intervals, row.finca, period, origin)
                data = intervals[(intervals.finca == row.finca) & (intervals.periodo == period) & (intervals.fecha <= origin)]
                posterior = DirichletM3(data, prior, cfg["bayes"]["dirichlet_prior_strength"], cfg["bayes"]["seed"])
                matrices = []
                for _ in range(cfg["bayes"]["posterior_draws"]):
                    q, r, loss = posterior.draw_matrix()
                    matrices.append(M3Matrix(row.finca, period, q, r, loss, prior.audit))
                posterior_cache[cache_key] = matrices
            x0 = np.array([row.RC_t0, row.SS_t0, row.AP_t0], float)
            lead = (pd.Timestamp(row.fecha_objetivo) - origin).days
            draws = [simulate(matrix, x0, lead, cfg["m3"]["baseline_ingress"]).iloc[-1].PC_dia_muestra * row.factor_extrapolacion
                     for matrix in posterior_cache[cache_key]]
            predictions.append(np.mean(draws))
        outputs["M3_DIRICHLET_MULTINOMIAL_ENTRADA"].append(_base_rows(
            frame, current, "M3_DIRICHLET_MULTINOMIAL_ENTRADA", predictions))
    return {key: pd.concat(value, ignore_index=True) if value else pd.DataFrame()
            for key, value in outputs.items()}


def _weekly(daily: pd.DataFrame) -> pd.DataFrame:
    if daily.empty:
        return daily
    keys = ["modelo", "finca", "bloque", "fecha_origen", "semana_proyeccion"]
    rows = []
    for key, group in daily.groupby(keys, sort=False, dropna=False):
        group = group.sort_values("fecha_objetivo")
        factor = (pd.to_numeric(group["factor_extrapolacion"], errors="coerce")
                  if "factor_extrapolacion" in group else pd.Series(dtype=float))
        observed = group[group["real"].notna()]
        target_dates = pd.to_datetime(group["fecha_objetivo"], errors="coerce").dropna()
        observed_dates = pd.to_datetime(observed["fecha_objetivo"], errors="coerce").dropna()
        horizons = (set(pd.to_numeric(group["horizonte_dia"], errors="coerce").dropna().astype(int))
                    if "horizonte_dia" in group else set(range(1, len(target_dates) + 1)))
        complete_dates = (
            len(observed_dates) == 7
            and observed_dates.nunique() == 7
            and (observed_dates.max() - observed_dates.min()).days == 6
            and horizons == set(range(1, 8))
        )
        if complete_dates:
            status = "COMPLETA"
        elif len(observed_dates):
            status = "PARCIAL"
        else:
            status = "NO EVALUABLE"
        rows.append(dict(zip(keys, key)) | {
            "real": observed["real"].sum(min_count=1),
            "proyectado": observed["proyectado"].sum(min_count=1),
            "dias": len(target_dates),
            "dias_reales": len(observed_dates),
            "estado_evaluacion": status,
            "factor_extrapolacion_min": factor.min() if not factor.empty else np.nan,
            "factor_extrapolacion_max": factor.max() if not factor.empty else np.nan,
            "factor_extrapolacion_n": factor.nunique() if not factor.empty else 0,
        })
    result = pd.DataFrame(rows)
    result["diferencia"] = result.proyectado - result.real
    result["error_abs"] = result.diferencia.abs()
    result["razon_proyectado_real"] = np.where(result.real.ne(0), result.proyectado / result.real, np.nan)
    result["desviacion_pct"] = result.razon_proyectado_real - 1
    result.loc[result.estado_evaluacion.eq("NO EVALUABLE"), ["real", "proyectado", "diferencia", "error_abs", "razon_proyectado_real", "desviacion_pct"]] = np.nan
    result["indicador"] = result.razon_proyectado_real.map(_status)
    return result


def _status(value: float | None) -> str:
    if value is None or not np.isfinite(value):
        return "SIN REAL"
    if 0.93 <= value <= 1.07:
        return "ACIERTO"
    if 0.90 <= value < 0.93 or 1.07 < value <= 1.10:
        return "CERCA"
    return "NO ACIERTO"


def _write_sheet(workbook: Workbook, name: str, frame: pd.DataFrame) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(max(1, len(frame.columns)))}{len(frame) + 1}"
    for index, column in enumerate(frame.columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(28, max(12, len(str(column)) + 2))
        if column in {"razon_proyectado_real", "desviacion_pct"}:
            for cell in sheet[get_column_letter(index)][1:]:
                cell.number_format = "0.0%"
    if "indicador" in frame.columns:
        column = get_column_letter(list(frame.columns).index("indicador") + 1)
        rng = f"{column}2:{column}{len(frame) + 1}"
        for label, color in (("ACIERTO", "C6EFCE"), ("CERCA", "FFEB9C"), ("NO ACIERTO", "FFC7CE")):
            sheet.conditional_formatting.add(rng, CellIsRule(
                operator="equal", formula=[f'"{label}"'], fill=PatternFill("solid", fgColor=color)))


def evaluate_input(root: Path, input_path: Path) -> tuple[Path, Path]:
    cfg = load_config(root / "config/pipeline.yaml")
    out = root / cfg["paths"]["outputs"]
    datasets, evaluation = out / "datasets", out / "evaluation"
    template = pd.read_parquet(datasets / "fact_bloque_dia.parquet")
    external_fact = _read_input(input_path, cfg, template, root / cfg["paths"]["raw"])
    missing_scale = external_fact["es_fecha_conteo"] & ~external_fact["escala_disponible"]
    scoring_fact = external_fact.loc[~missing_scale].copy()
    windows = build_forecast_windows(scoring_fact, int(cfg["forecast"]["horizon_days"]))
    intervals = load_traditional_intervals(root / cfg["paths"]["raw"], cfg)
    pruning = pd.read_parquet(datasets / "poda_features.parquet")
    climate = pd.read_parquet(datasets / "clima_features.parquet")
    frame = build_supervised_dataset(windows, scoring_fact, intervals, cfg, pruning, climate, include_incomplete=True)
    historical = pd.read_parquet(datasets / "dataset_supervisado_diario.parquet")
    predictions = {
        "E00_M3_BASE_ENTRADA": _m3(frame, intervals, cfg, cfg["m3"]["baseline_ingress"], "E00_M3_BASE_ENTRADA"),
        "E01_M3_INGRESO_CALIBRADO_ENTRADA": _m3(frame, intervals, cfg, cfg["m3"]["calibrated_ingress"], "E01_M3_INGRESO_CALIBRADO_ENTRADA"),
    }
    best = pd.read_csv(evaluation / "rf_mejores_por_grupo.csv")
    rf_specs = {f"RF_OPT_{row.features}_ENTRADA": {"family": "RF", "features": row.features,
        "hyperparameters": json.loads(row.hyperparameters)} for _, row in best.iterrows()}
    predictions.update(_direct_models(frame, historical, cfg, rf_specs))
    feno_spec = rf_specs.get("RF_OPT_FENO_ENTRADA")
    if feno_spec:
        predictions.update(_additional_rf_models(frame, historical, cfg, feno_spec))
    predictions.update(_bayes_models(frame, historical, intervals, cfg))
    daily = pd.concat([value.assign(fuente_entrada=input_path.name) for value in predictions.values() if not value.empty], ignore_index=True)
    weekly = _weekly(daily)
    evaluation.mkdir(parents=True, exist_ok=True)
    daily_path = evaluation / "predictions_entrada_diarias.csv"
    weekly_path = evaluation / "predictions_entrada_semanales.csv"
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "SEMANAL", weekly)
    _write_sheet(workbook, "DIARIO", daily)
    workbook.save(evaluation / "evaluacion_entrada.xlsx")
    daily.to_csv(daily_path, index=False)
    weekly.to_csv(weekly_path, index=False)
    factor_audit = (daily.groupby(["modelo", "finca", "bloque", "semana_proyeccion"], as_index=False)
                    .agg(filas=("factor_extrapolacion", "size"),
                         factor_min=("factor_extrapolacion", "min"),
                         factor_max=("factor_extrapolacion", "max"),
                         factor_n=("factor_extrapolacion", "nunique"),
                         factores_faltantes=("factor_extrapolacion", lambda values: int(values.isna().sum()))))
    factor_audit_path = evaluation / "factor_extrapolacion_auditoria.csv"
    factor_audit.to_csv(factor_audit_path, index=False)
    manifest = {
        "population": "EXTERNAL_SCORING",
        "input_file": str(input_path.relative_to(root)) if input_path.is_relative_to(root) else str(input_path),
        "input_sha256": hashlib.sha256(input_path.read_bytes()).hexdigest(),
        "historical_dataset": str((datasets / "dataset_supervisado_diario.parquet").relative_to(root)),
        "historical_max_target_date": pd.to_datetime(historical["fecha_objetivo"]).max().strftime("%Y-%m-%d"),
        "historical_training_rows": int(historical["target"].notna().sum()),
        "models": sorted(daily["modelo"].unique().tolist()),
        "daily_rows": int(len(daily)),
        "weekly_rows": int(len(weekly)),
        "partial_week_rows": int(weekly["estado_evaluacion"].eq("PARCIAL").sum()),
        "complete_week_rows": int(weekly["estado_evaluacion"].eq("COMPLETA").sum()),
        "input_rows_without_extrapolation": int(missing_scale.sum()),
        "input_keys_without_extrapolation": int(
            external_fact.loc[missing_scale, ["finca", "bloque", "semana_iso"]]
            .drop_duplicates().shape[0]),
        "extrapolation_policy": "exclude_origin_without_valid_scale",
        "factor_audit_file": str(factor_audit_path.relative_to(root)),
        "factor_audit_rows": int(len(factor_audit)),
        "factor_audit_missing": int(factor_audit.factores_faltantes.sum()),
    }
    (evaluation / "external_run_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    return daily_path, weekly_path


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[1]
    input_path = Path(sys.argv[1]) if len(sys.argv) > 1 else (
        root / load_config(root / "config/pipeline.yaml")["paths"]["external"] /
        load_config(root / "config/pipeline.yaml")["vision"]["operational_history_path"])
    print("\n".join(map(str, evaluate_input(root, input_path))))
